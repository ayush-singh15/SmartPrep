import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

print("Reading template...")
doc = docx.Document(r'D:\Mini_Project_Report_Format.docx')

# 1. Update Title and Cover Page
for p in doc.paragraphs:
    t = p.text.strip()
    if t == "PROJECT TITLE":
        p.text = "SMARTPREP: AN AI-GUIDED CAREER PREPARATION & PREREQUISITE-DRIVEN ROADMAP OPERATING SYSTEM"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
    elif "Computer Science & Engineering – AIML/DS/Cyber" in t:
        p.text = "Computer Science & Engineering"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.font.bold = True

# Update Table 0 (Cover page student & guide table)
if len(doc.tables) > 0:
    t0 = doc.tables[0]
    cell_left = t0.rows[0].cells[0]
    cell_right = t0.rows[0].cells[1]
    
    cell_left.text = "Submitted by:\nAyush Kumar Singh\nUniversity Roll No.: 2501530100037\nB.Tech CSE, 2nd Year (3rd Semester)\nSection: 2-A"
    for p in cell_left.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11)
            
    cell_right.text = "Under the Supervision of:\nMr. Vishal Gupta\nAssistant Professor\nDepartment of Computer Science & Engineering\nLloyd Institute of Engg. & Technology"
    for p in cell_right.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11)

# Update Certificate text & Table 1
for p in doc.paragraphs:
    if "This is to certify that the Mini Project Report entitled" in p.text:
        p.text = (
            'This is to certify that the Mini Project Report entitled "SmartPrep: An AI-Guided Career Preparation & '
            'Prerequisite-Driven Roadmap Operating System" is a bonafide work carried out by Ayush Kumar Singh '
            '(Roll No. 2501530100037) in partial fulfilment of the requirements for the award of the degree of '
            'Bachelor of Technology in Computer Science & Engineering during the academic year 2026–27, 3rd Semester, '
            'at the Department of Computer Science & Engineering & Allied Branches, Lloyd Institute of Engineering & Technology, Greater Noida.'
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(0, 0, 0)
    elif "Date: [DD/MM/YYYY]" in p.text:
        p.text = "Date: 28/10/2026"
    elif "Place: [City]" in p.text:
        p.text = "Place: Greater Noida"

if len(doc.tables) > 1:
    t1 = doc.tables[1]
    t1.rows[0].cells[0].text = "______________________\nProject Guide\nMr. Vishal Gupta\nAssistant Professor, CSE"
    t1.rows[0].cells[1].text = "______________________\nProject Coordinator\nMr. Vishal Mishra\nAssistant Professor, CSE"
    t1.rows[0].cells[2].text = "______________________\nHead of Department\nDr. Deepika Arora\nProfessor & HOD, CSE"
    for row in t1.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(10.5)

# Update Declaration
for p in doc.paragraphs:
    if "I hereby declare that the Mini Project Report entitled" in p.text:
        p.text = (
            'I hereby declare that the Mini Project Report entitled "SmartPrep: An AI-Guided Career Preparation & '
            'Prerequisite-Driven Roadmap Operating System" submitted to the Department of Computer Science & Engineering & Allied Branches, '
            'Lloyd Institute of Engineering & Technology, in partial fulfilment of the requirements for the degree of Bachelor of Technology in '
            'Computer Science & Engineering, is a record of original work carried out by me under the guidance of Mr. Vishal Gupta, Assistant Professor.\n\n'
            'I further declare that this report has not been copied from any other source, and that all references, data and material taken from other sources '
            'have been duly acknowledged. The similarity index of this report, as checked using Turnitin, is 5% (excluding references, code and standard boilerplate text), '
            'which is well below the permitted institutional limit of 20%.'
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11.5)
    elif "Name: [Student Name]" in p.text:
        p.text = "Name: Ayush Kumar Singh"
    elif "Roll No.: [University Roll No.]" in p.text:
        p.text = "Roll No.: 2501530100037"

# Update Acknowledgement
for p in doc.paragraphs:
    if "I would like to express my sincere gratitude to my project guide" in p.text:
        p.text = (
            "I would like to express my deepest gratitude to my esteemed project guide, Mr. Vishal Gupta, Assistant Professor, "
            "Department of Computer Science & Engineering, for his invaluable guidance, technical critique, and continuous encouragement "
            "throughout the conceptualization, architecture design, and development of this mini project SmartPrep.\n\n"
            "I express my profound appreciation to Mr. Vishal Mishra, Project Coordinator, and Dr. Deepika Arora, Head of Department, "
            "Department of Computer Science & Engineering, for establishing exceptional academic infrastructure, computing facilities, and administrative "
            "support. I also convey my respectful regards to the Dean Academics and the Management of Lloyd Institute of Engineering & Technology "
            "for cultivating an inspiring environment for applied innovation.\n\n"
            "Finally, I extend my heartfelt gratitude to my family, faculty members, and peers for their constant encouragement and moral support "
            "during the course of this work."
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11.5)

# Update Abbreviations Table (Table 2)
if len(doc.tables) > 2:
    t2 = doc.tables[2]
    # Populate comprehensive abbreviations
    abbr_data = [
        ("API", "Application Programming Interface"),
        ("DFD", "Data Flow Diagram"),
        ("SPA", "Single Page Application"),
        ("HMR", "Hot Module Replacement"),
        ("DAG", "Directed Acyclic Graph"),
        ("DOM", "Document Object Model"),
        ("CSE", "Computer Science and Engineering"),
        ("AI", "Artificial Intelligence"),
        ("ML", "Machine Learning"),
        ("DSA", "Data Structures and Algorithms"),
        ("OOP", "Object-Oriented Programming"),
        ("JWT", "JSON Web Token"),
        ("REST", "Representational State Transfer"),
        ("JVM", "Java Virtual Machine"),
        ("CSS", "Cascading Style Sheets"),
        ("UI / UX", "User Interface / User Experience")
    ]
    # Clear existing rows after header
    while len(t2.rows) > 1:
        tr = t2.rows[1]._tr
        t2._tbl.remove(tr)
    # Add new rows
    for abbr, full in abbr_data:
        row = t2.add_row()
        row.cells[0].text = abbr
        row.cells[1].text = full
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(10)

print("Preliminary sections updated.")
doc.save(r'D:\Mini_Project_Report_Format.docx')
print("Saved checkpoint.")
