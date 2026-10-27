import pptx

prs = pptx.Presentation(r'D:\Mini_Project_Review-02_Final_Evaluation - Format.pptx')
with open('slides_deep.txt', 'w', encoding='utf-8') as out_f:
    for idx, slide in enumerate(prs.slides, 1):
        out_f.write(f"==================== SLIDE {idx} ====================\n")
        for sh_idx, shape in enumerate(slide.shapes):
            txt = ""
            if shape.has_text_frame:
                txt = shape.text_frame.text.replace("\n", " \\n ")
            tbl = ""
            if shape.has_table:
                tbl = f"TABLE {len(shape.table.rows)}x{len(shape.table.columns)}"
            out_f.write(f"  Shape {sh_idx}: id={shape.shape_id}, name='{shape.name}', type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height} | {txt[:100]} {tbl}\n")
print("Done writing slides_deep.txt")
